"""규칙 기반 간이 추출기 (데모용 기준 추출기). devplan 8.1~8.2.

문서(xlsx, docx)에서 요구사항을 정규식과 라벨 일치로 뽑아 채점 입력 형식(docs/specs/GRADE_INPUT_SPEC.md)으로 낸다.

**한계 (읽기 전에)**
- 이것은 데모용이다. **규칙 기반 점수는 LLM 추출의 성능을 대표하지 않는다.** 채점기가 어디서 왜 틀리는지를 보여 주는
  재료일 뿐이다.
- 추출기는 문서와 스키마의 공개 정보(항목 이름, 단위, 별칭)만 쓴다. 정답(ground_truth)과 표시 계획(document_plan),
  렌더러 내부는 보지 않는다. 문서를 만든 양식의 이름이나 배치도 코드에 적지 않는다.
- **양식 A·B로만 개발·튜닝했다.** 보류 양식 C는 최종 평가 때 한 번만 보았다. 튜닝 이력은 docs/decisions.md에 있다.

동작 방식
  엑셀: 모든 시트의 모든 행에서 항목 이름과 같은 칸을 찾고, 그 오른쪽 칸들에서 값과 단위를 읽는다.
  워드: 모든 문단(표 안 포함)이 항목 이름으로 시작하면, 그 뒤 문장에서 값과 단위를 읽는다.
  단위·선택 값은 스키마가 아는 표기(별칭 포함)와 앞부분이 일치하면 그 표기로, 아니면 보이는 그대로 낸다.
  알 수 없는 표기(오타 등)는 고치지 않고 그대로 내서 채점기가 형식 오류로 세게 한다.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

_SRC = Path(__file__).resolve().parent.parent / "src"
if _SRC.is_dir() and str(_SRC) not in sys.path:  # 저장소에서 바로 실행할 때 synkit을 찾게 한다
    sys.path.insert(0, str(_SRC))

import docx  # noqa: E402
import openpyxl  # noqa: E402

from synkit.schema.labels import label_for  # noqa: E402
from synkit.schema.loader import load_schema_file  # noqa: E402
from synkit.schema.models import Requirement  # noqa: E402

_NUMBER = re.compile(r"-?\d+(?:\.\d+)?")
_LIST_SEPARATOR = re.compile(r"[,，、;]")
_LEADING = re.compile(r"^\s*(?:[:：]|은|는|이|가)?\s*")


# ---------------------------------------------------------------- 스키마의 공개 정보

def _key(text) -> str:
    return " ".join(str(text).casefold().split())


def _label_map(requirements: list[Requirement]) -> dict[str, Requirement]:
    mapping = {}
    for requirement in requirements:
        for language in ("ko", "en"):
            mapping[_key(label_for(requirement, language))] = requirement
    return mapping


def _unit_spellings(requirement: Requirement) -> list[str]:
    aliases = requirement.unit_aliases or {}
    spellings = [requirement.unit] if requirement.unit else []
    spellings += aliases.get(requirement.unit, [])
    for source in requirement.conversions or {}:
        spellings.append(source)
        spellings += aliases.get(source, [])
    return sorted({s for s in spellings if s}, key=len, reverse=True)


def _value_spellings(requirement: Requirement) -> list[str]:
    spellings = list(requirement.values or [])
    for alias_list in (requirement.aliases or {}).values():
        spellings += alias_list
    return sorted(set(spellings), key=len, reverse=True)


_WORD_CHAR = re.compile(r"[A-Za-z0-9_]")


def _known_prefix(text: str, spellings: list[str]) -> str | None:
    """text가 스키마가 아는 표기로 시작하면 그 부분을 돌려준다.

    영문자·숫자로 끝나는 표기는 바로 뒤에 영문자·숫자가 오면 같은 표기가 아니다. (`volst`는 `V`가 아니다)
    한글 조사가 바로 붙는 것(`12V이어야`)은 허용한다. 이 경계가 없으면 오타가 알려진 표기로 몰래 고쳐져 점수가 부풀려진다.
    """
    folded = text.casefold()
    for spelling in spellings:
        if not folded.startswith(spelling.casefold()):
            continue
        following = text[len(spelling):len(spelling) + 1]
        if _WORD_CHAR.match(spelling[-1]) and following and _WORD_CHAR.match(following):
            continue
        return text[: len(spelling)]
    return None


# ---------------------------------------------------------------- 값 읽기

def _to_number(text: str):
    return int(text) if re.fullmatch(r"-?\d+", text) else float(text)


def _split_list(text: str) -> list[str]:
    return [part.strip().rstrip(".。 ") for part in _LIST_SEPARATOR.split(text) if part.strip().rstrip(".。 ")]


def _read_unit(after: str, requirement: Requirement) -> str | None:
    after = after.strip()
    if not after:
        return None
    known = _known_prefix(after, _unit_spellings(requirement))
    if known:
        return known
    return after.split()[0].rstrip(".,;:") or None


def _number_entry(requirement: Requirement, text: str, next_cell=None):
    """'12 V', '350 m3/h', '이어야 12 V ...'처럼 숫자와 단위가 한 문자열에 있는 경우."""
    match = _NUMBER.search(text)
    if not match:
        return {"id": requirement.id, "value": text.strip() or None, "raw": text}
    unit = _read_unit(text[match.end():], requirement)
    if unit is None and isinstance(next_cell, str):
        unit = next_cell.strip() or None
    entry = {"id": requirement.id, "value": _to_number(match.group()), "raw": text.strip()}
    if unit is not None:
        entry["unit"] = unit
    return entry


def _choice_entry(requirement: Requirement, text: str):
    text = text.strip().rstrip(".。 ")
    if requirement.type == "multi_select":
        return {"id": requirement.id, "value": _split_list(text), "raw": text}
    known = _known_prefix(text, _value_spellings(requirement)) if requirement.type == "category" else None
    return {"id": requirement.id, "value": known or text, "raw": text}


# ---------------------------------------------------------------- 엑셀

def _from_cells(requirement: Requirement, cells: list):
    first = cells[0]
    if requirement.type == "number":
        if isinstance(first, (int, float)) and not isinstance(first, bool):
            unit = cells[1].strip() if len(cells) > 1 and isinstance(cells[1], str) and cells[1].strip() else None
            entry = {"id": requirement.id, "value": first, "raw": f"{first} {unit or ''}".strip()}
            if unit is not None:
                entry["unit"] = unit
            return entry
        return _number_entry(requirement, str(first), cells[1] if len(cells) > 1 else None)
    return _choice_entry(requirement, str(first)) if requirement.type != "text" else {
        "id": requirement.id, "value": str(first).strip(), "raw": str(first)}


def _extract_excel(path: Path, labels: dict[str, Requirement]) -> dict[str, dict]:
    found: dict[str, dict] = {}
    for sheet in openpyxl.load_workbook(path).worksheets:
        for row in sheet.iter_rows(values_only=True):
            cells = list(row)
            for index, cell in enumerate(cells):
                if isinstance(cell, str) and _key(cell) in labels:
                    requirement = labels[_key(cell)]
                    rest = [c for c in cells[index + 1:] if c is not None and not (isinstance(c, str) and not c.strip())]
                    if rest and requirement.id not in found:
                        found[requirement.id] = _from_cells(requirement, rest)
                    break
    return found


# ---------------------------------------------------------------- 워드

def _paragraphs(document) -> list[str]:
    texts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                texts += [p.text for p in cell.paragraphs]
    return [t.strip() for t in texts if t.strip()]


def _extract_word(path: Path, labels: dict[str, Requirement]) -> dict[str, dict]:
    found: dict[str, dict] = {}
    ordered = sorted(labels, key=len, reverse=True)  # 긴 항목 이름부터 맞춰 본다
    for text in _paragraphs(docx.Document(str(path))):
        for label in ordered:
            if _key(text).startswith(label):
                requirement = labels[label]
                remainder = text[len(label):]
                if requirement.id not in found:
                    if requirement.type == "number":
                        found[requirement.id] = _number_entry(requirement, remainder)
                    elif requirement.type == "text":
                        found[requirement.id] = {"id": requirement.id, "value": _LEADING.sub("", remainder).strip().rstrip(".。 "), "raw": text}
                    else:
                        found[requirement.id] = _choice_entry(requirement, _LEADING.sub("", remainder))
                break
    return found


# ---------------------------------------------------------------- 공개 함수

def extract_document(path: str | Path, requirements: list[Requirement]) -> list[dict]:
    """문서 하나에서 모든 요구사항을 뽑는다. 찾지 못한 항목은 value가 None이다."""
    path = Path(path)
    labels = _label_map(requirements)
    if path.suffix == ".xlsx":
        found = _extract_excel(path, labels)
    elif path.suffix == ".docx":
        found = _extract_word(path, labels)
    else:
        raise ValueError(f"지원하지 않는 문서 형식입니다: {path.suffix!r}")
    return [found.get(r.id, {"id": r.id, "value": None}) for r in requirements]


def extract_run(run_path: str | Path) -> dict:
    """실행 폴더의 documents/를 모두 읽어 채점 입력 형식의 dict로 낸다."""
    run_path = Path(run_path)
    loaded = load_schema_file(run_path / "schema.yaml")
    if not loaded.ok:
        raise ValueError("실행 폴더의 스키마 사본을 읽을 수 없습니다.")
    meta = json.loads((run_path / "run_meta.json").read_text(encoding="utf-8"))
    requirements = loaded.schema.requirements
    paths = sorted(p for p in (run_path / "documents").iterdir() if p.suffix in (".xlsx", ".docx"))
    return {
        "schema_version": "1",
        "run_id": meta["run_id"],
        "results": [{"doc_id": p.stem, "requirements": extract_document(p, requirements)} for p in paths],
    }


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="규칙 기반 간이 추출기 (데모용)")
    parser.add_argument("run_dir", type=Path, help="render까지 끝낸 실행 폴더")
    parser.add_argument("out", type=Path, help="추출 결과를 쓸 JSON 파일")
    args = parser.parse_args(argv)
    data = extract_run(args.run_dir)
    args.out.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"추출 완료: 문서 {len(data['results'])}건 → {args.out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
